from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
import numpy_financial as npf
import pandas as pd
import io
from typing import List
from app.database import get_db
from app.models import Unit, Simulation, SimulationResult, SimulationDetail, Client, Advisor, Prospect, BonoBBP, CreditoIFI, User, ModalidadVivienda, TipoVenta
from app.schemas.simulation import SimulationCreate, SimulationResponse
from app.core.security import get_current_user
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import openpyxl

router = APIRouter()

def d2(val):
    if isinstance(val, (float, int)):
        val = Decimal(str(val))
    return val.quantize(Decimal('0.00'), rounding=ROUND_HALF_UP)

def get_bono_info(db: Session, precio_venta: Decimal, tipo_bbp: str, modalidad: str) -> dict:
    """Calcula rango, bono base e integrador según tablas técnicas."""
    if modalidad == "Mejoramiento" or precio_venta > Decimal("488800.00") or tipo_bbp == "Ninguno":
        return {"rango": "SinBBP", "base": Decimal("0.00"), "integrador": Decimal("0.00")}

    bono_row = db.query(BonoBBP).filter(
        BonoBBP.valor_vivienda_min <= precio_venta,
        BonoBBP.valor_vivienda_max >= precio_venta
    ).first()

    if not bono_row or bono_row.rango == "R5":
        return {"rango": (bono_row.rango if bono_row else "SinBBP"), "base": Decimal("0.00"), "integrador": Decimal("0.00")}

    mapping = {
        "Tradicional": bono_row.bono_tradicional,
        "Sostenible": bono_row.bono_sostenible,
        "Integrador Tradicional": bono_row.bono_tradicional,
        "Integrador Sostenible": bono_row.bono_sostenible,
    }
    
    base = Decimal(str(mapping.get(tipo_bbp, 0)))
    integrador = Decimal("3600.00") if "Integrador" in tipo_bbp else Decimal("0.00")
    
    return {"rango": bono_row.rango, "base": base, "integrador": integrador}

def get_capitalizacion_factor(capitalizacion: str) -> int:
    """Retorna el número de meses de capitalización para tasa nominal."""
    return {"Mensual": 1, "Bimestral": 2, "Trimestral": 3}.get(capitalizacion, 1)

@router.get("/bbp-ranges")
def get_bbp_ranges(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """
    Devuelve los rangos de bonos BBP configurados en el sistema.
    """
    bonos = db.query(BonoBBP).order_by(BonoBBP.valor_vivienda_min.asc()).all()
    return [
        {
            "rango": b.rango,
            "min": float(b.valor_vivienda_min),
            "max": float(b.valor_vivienda_max),
            "tradicional": float(b.bono_tradicional),
            "sostenible": float(b.bono_sostenible),
            "integrador_tradicional": float(b.bono_integrador_tradicional),
            "integrador_sostenible": float(b.bono_integrador_sostenible)
        }
        for b in bonos
    ]

# 🚨 NUEVO ENDPOINT: Búsqueda segura del IFM para el Frontend
@router.get("/check-income/{person_id}")
def check_income(person_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    # Busca primero como prospecto
    prospect = db.query(Prospect).filter(Prospect.codigo_prospecto == person_id).first()
    if prospect:
        return {
            "type": "Prospecto", 
            "ifm": float(prospect.ingreso_mensual or 0) + float(getattr(prospect, 'ingreso_conyuge', 0))
        }
    
    # Si no, busca como cliente
    client = db.query(Client).filter(Client.codigo_cliente == person_id).first()
    if client:
        return {
            "type": "Cliente", 
            "ifm": float(client.ingreso_mensual or 0) + float(getattr(client, 'ingreso_conyuge', 0))
        }
        
    raise HTTPException(status_code=404, detail="ID no encontrado en Base de Datos")

@router.get("/ifis-disponibles")
def get_ifis_disponibles(
    monto: float,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Devuelve las IFIs disponibles para un monto de financiamiento dado.
    """
    from decimal import Decimal
    monto_dec = Decimal(str(monto))
    results = db.query(CreditoIFI).filter(
        CreditoIFI.monto_min <= monto_dec,
        (CreditoIFI.monto_max >= monto_dec) | (CreditoIFI.monto_max == None)
    ).all()

    return [
        {
            "nombre_ifi": r.nombre_ifi,
            "plazo_min_anios": r.plazo_min_anios,
            "plazo_max_anios": r.plazo_max_anios,
            "tea_min": float(r.tea_min),
            "tea_max": float(r.tea_max),
            "seguro_individual": float(r.seguro_individual),
            "seguro_mancomunado": float(r.seguro_mancomunado),
        }
        for r in results
    ]


@router.post("/", response_model=SimulationResponse)
def run_simulation(
    payload: SimulationCreate,
    save: bool = Query(False),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # 1. Validar unidad y obtener modalidad
    unit = db.query(Unit).filter(Unit.codigo_unidad == payload.codigo_unidad).first()
    if not unit:
        raise HTTPException(status_code=404, detail="Unidad no encontrada")

    modalidad = "Compra"
    if unit.codigo_modalidad:
        mod = db.query(ModalidadVivienda).filter(ModalidadVivienda.codigo_modalidad == unit.codigo_modalidad).first()
        modalidad = mod.nombre_modalidad if mod else "Compra"

    role = current_user.rol_rel.tipo_rol
    
    # 2. Asignar identificadores y buscar entidad para IFM
    entity = None
    ifm = Decimal("0")

    if role == "Cliente":
        if payload.codigo_cliente and payload.codigo_cliente != current_user.codigo_usuario:
            raise HTTPException(status_code=403, detail="No puedes crear simulaciones en nombre de otro cliente.")
        payload.codigo_cliente = current_user.codigo_usuario
        payload.codigo_asesor = None
        payload.codigo_prospecto = None
        
        entity = db.query(Client).filter(Client.codigo_cliente == current_user.codigo_usuario).first()
        if entity: 
            ifm = Decimal(str(entity.ingreso_mensual or 0)) + Decimal(str(getattr(entity, 'ingreso_conyuge', 0)))

    elif role == "Asesor":
        payload.codigo_asesor = current_user.codigo_usuario
        input_id = payload.codigo_prospecto
        
        if not input_id:
            raise HTTPException(status_code=400, detail="El asesor debe indicar el ID del prospecto/cliente.")

        is_prospect = db.query(Prospect).filter(Prospect.codigo_prospecto == input_id).first()
        is_client = db.query(Client).filter(Client.codigo_cliente == input_id).first()
        
        if is_prospect:
            payload.codigo_prospecto = input_id
            payload.codigo_cliente = None
            entity = is_prospect
        elif is_client:
            payload.codigo_cliente = input_id
            payload.codigo_prospecto = None  
            entity = is_client
        else:
            raise HTTPException(status_code=404, detail=f"El ID {input_id} no existe en la base de datos.")

        ifm = Decimal(str(entity.ingreso_mensual or 0)) + Decimal(str(getattr(entity, 'ingreso_conyuge', 0)))

    TIPOS_BBP = ["Ninguno", "Tradicional", "Sostenible", "Integrador Tradicional", "Integrador Sostenible"]
    if payload.tipo_bbp not in TIPOS_BBP:
        raise HTTPException(status_code=422, detail=f"tipo_bbp debe ser uno de: {TIPOS_BBP}")
    
    CATEGORIAS_INTEGRADOR = ["Menores ingresos", "Adulto mayor", "Discapacidad", "Desplazado", "Migrante retornado"]
    if "Integrador" in payload.tipo_bbp:
        if not payload.categoria_integrador:
            raise HTTPException(status_code=422, detail="categoria_integrador es requerida cuando tipo_bbp es Integrador.")
        if payload.categoria_integrador not in CATEGORIAS_INTEGRADOR:
            raise HTTPException(status_code=422, detail=f"categoria_integrador debe ser una de: {CATEGORIAS_INTEGRADOR}")
        if payload.categoria_integrador == "Menores ingresos":
            if payload.ingreso_maximo_integrador is None:
                raise HTTPException(status_code=422, detail="ingreso_maximo_integrador es requerido para 'Menores ingresos'.")
            if payload.ingreso_maximo_integrador > 4746.00:
                raise HTTPException(status_code=422, detail="ingreso_maximo_integrador no puede superar S/ 4,746.00")
    else:
        payload.categoria_integrador = None
        payload.ingreso_maximo_integrador = None

    if payload.meses_gracia >= payload.plazo_meses:
        raise HTTPException(status_code=422, detail="meses_gracia debe ser menor que plazo_meses.")

    if payload.tipo_gracia == "Total" and payload.meses_gracia > 6:
        raise HTTPException(status_code=422, detail="El periodo de gracia total no puede superar los 6 meses.")

    # 👇 ESTA ES LA LÍNEA MÁGICA QUE FALTABA (YA AGREGADA Y CORREGIDA) 👇
    fecha_base = payload.fecha_inicio_prestamo if payload.fecha_inicio_prestamo else date.today()

    # 3. Conversión de Moneda y Cálculo de BBP
    pv = Decimal(str(unit.precio_venta))
    tipo_cambio = Decimal(str(payload.tipo_cambio))

    pv_pen = pv
    if unit.moneda_rel.simbolo_moneda == "USD":
        pv_pen = pv * tipo_cambio

    bono_info = get_bono_info(db, pv_pen, payload.tipo_bbp, modalidad)
    rango = bono_info["rango"]
    
    bono_base = bono_info["base"]
    bono_integrador = bono_info["integrador"]
    if unit.moneda_rel.simbolo_moneda == "USD":
        bono_base = (bono_base / tipo_cambio).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        bono_integrador = (bono_integrador / tipo_cambio).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    
    bono_total = bono_base + bono_integrador

    if payload.tipo_bbp != "Ninguno" and not payload.ifi_seleccionada:
        raise HTTPException(status_code=400, detail="El BBP solo aplica con el Crédito MiVivienda a través de una IFI.")

    if unit.es_sostenible:
        if payload.tipo_bbp in ["Tradicional", "Integrador Tradicional"]:
            raise HTTPException(status_code=400, detail="Vivienda Sostenible solo acepta bonos Sostenibles.")
    else:
        if payload.tipo_bbp in ["Sostenible", "Integrador Sostenible"]:
            raise HTTPException(status_code=400, detail="Vivienda Tradicional solo acepta bonos Tradicionales.")

    if payload.tipo_bbp != "Ninguno":
        tipo_v = db.query(TipoVenta).filter(TipoVenta.codigo_tipo_venta == unit.codigo_tipo_venta).first()
        if tipo_v and tipo_v.nombre_tipo_venta == "Segunda venta":
            raise HTTPException(status_code=400, detail="El Bono de Buen Pagador (BBP) solo aplica para unidades de Primera Venta.")

    if payload.tipo_bbp != "Ninguno" and entity:
        if hasattr(entity, 'es_propietario_vivienda') and entity.es_propietario_vivienda:
            raise HTTPException(status_code=400, detail="El solicitante ya es propietario de una vivienda. No califica para el BBP.")
        
        if hasattr(entity, 'conyuge_propietario') and entity.conyuge_propietario:
            raise HTTPException(status_code=400, detail="El cónyuge o conviviente ya es propietario de una vivienda. No califica para el BBP.")
        elif hasattr(entity, 'es_casado') and entity.es_casado and hasattr(entity, 'conyuge_rel'):
            pass

        if hasattr(entity, 'hijos_menores_propietarios') and entity.hijos_menores_propietarios:
            raise HTTPException(status_code=400, detail="Uno o más hijos menores de edad figura como propietario. No califica para el BBP.")
        
        if hasattr(entity, 'recibio_apoyo_estatal') and entity.recibio_apoyo_estatal:
            raise HTTPException(status_code=400, detail="El solicitante ya recibió apoyo habitacional del Estado. No califica para el BBP.")
        
        if hasattr(entity, 'cantidad_creditos_fmv') and entity.cantidad_creditos_fmv >= 2:
            raise HTTPException(status_code=400, detail="Ha alcanzado el límite máximo de 2 créditos MiVivienda.")
        
        if hasattr(entity, 'tiene_credito_fmv_activo') and entity.tiene_credito_fmv_activo:
            raise HTTPException(status_code=400, detail="El solicitante tiene un Crédito MiVivienda activo. Bloqueado.")

    # 4. Cuota Inicial y Financiamiento
    cuota_inicial = Decimal(str(payload.cuota_inicial))
    porcentaje_inicial = (cuota_inicial / pv) * 100
    min_porcentaje = Decimal("7.5") if modalidad in ["Construccion", "Mejoramiento"] else Decimal("10.0")

    if porcentaje_inicial < min_porcentaje:
        raise HTTPException(status_code=400, detail=f"La cuota inicial es insuficiente. Se requiere un mínimo del {float(min_porcentaje)}%.")

    # 5. Cálculo del Préstamo
    precio_neto = pv - bono_total
    monto_financiar = (precio_neto - cuota_inicial) + Decimal(str(payload.gastos_iniciales))

    if modalidad not in ["Construccion", "Mejoramiento"]:
        max_financiar = pv * Decimal("0.90")
        if monto_financiar > max_financiar:
            raise HTTPException(status_code=400, detail=f"Financiamiento excede el 90% (S/ {float(max_financiar):,.2f}).")

    max_gastos = pv * Decimal("0.05")
    gastos_actuales = Decimal(str(payload.gastos_iniciales or 0))
    if gastos_actuales > max_gastos:
        raise HTTPException(
            status_code=400, 
            detail=f"Los gastos iniciales (S/ {float(gastos_actuales):,.2f}) exceden el límite del 5% (Máx S/ {float(max_gastos):,.2f})."
        )

    if monto_financiar <= 0:
        raise HTTPException(status_code=400, detail="El monto a financiar debe ser mayor a 0.")

    if payload.ifi_seleccionada:
        ifi_row = db.query(CreditoIFI).filter(
            CreditoIFI.nombre_ifi == payload.ifi_seleccionada,
            CreditoIFI.monto_min <= monto_financiar,
            (CreditoIFI.monto_max >= monto_financiar) | (CreditoIFI.monto_max == None)
        ).first()

        if not ifi_row:
            raise HTTPException(status_code=400, detail=f"Monto fuera de rango para {payload.ifi_seleccionada}.")

        payload.tipo_tasa = "Efectiva"
        payload.capitalizacion = "Mensual" 

        tasa_ingresada = Decimal(str(payload.tasa_anual))
        if not (ifi_row.tea_min <= tasa_ingresada <= ifi_row.tea_max):
            raise HTTPException(status_code=400, detail=f"La tasa está fuera del rango TEA de {payload.ifi_seleccionada}.")
        
        plazo_anios = payload.plazo_meses / 12
        if not (ifi_row.plazo_min_anios <= plazo_anios <= ifi_row.plazo_max_anios):
            raise HTTPException(status_code=400, detail=f"El plazo está fuera del rango de {payload.ifi_seleccionada}.")
        
        tiene_mancomunado = getattr(entity, 'tiene_deudor_solidario', False) if entity else False
        payload.seguro_desgravamen = float(ifi_row.seguro_mancomunado if tiene_mancomunado else ifi_row.seguro_individual)

    else:
        if payload.tipo_tasa == "Efectiva":
            payload.capitalizacion = "Mensual"

    tasa_anual_dec = Decimal(str(payload.tasa_anual)) / Decimal("100")
    if payload.tipo_tasa == "Efectiva":
        tea = tasa_anual_dec
    else: 
        m = Decimal("12") / Decimal(str({"Mensual": 1, "Bimestral": 2, "Trimestral": 3}.get(payload.capitalizacion, 1)))
        tea = (1 + tasa_anual_dec / m)**m - 1
    
    tem = (1 + tea)**(Decimal("1")/Decimal("12")) - 1
    
    seguro_tasa = Decimal(str(payload.seguro_desgravamen)) / Decimal("100")
    n_total = payload.plazo_meses
    m_gracia = payload.meses_gracia
    n_reales = n_total - m_gracia
    
    seguro_tasa = Decimal(str(payload.seguro_desgravamen)) / Decimal("100")
    tasa_para_factor = tem + seguro_tasa
    
    factor = (tasa_para_factor * (1 + tasa_para_factor) ** n_reales) / ((1 + tasa_para_factor) ** n_reales - 1) if tasa_para_factor > 0 else (Decimal("1")/Decimal(str(n_reales)))
    cuota_base = monto_financiar * factor

    flujos_caja = [float(monto_financiar) - float(payload.gastos_iniciales)]
    total_int, total_seg = Decimal("0"), Decimal("0")

    detalles_db = []
    saldo = monto_financiar
    
    flujo_0 = monto_financiar - Decimal(str(payload.gastos_iniciales))
    detalles_db.append(SimulationDetail(
        numero_cuota=0,
        saldo_inicio=d2(monto_financiar),
        interes=d2(0),
        interes_capitalizado=d2(0),
        seguro=d2(0),
        amortizacion=d2(0),
        cuota_total=d2(0),
        saldo_final=d2(monto_financiar),
        flujo_caja=d2(flujo_0),
        fecha_vencimiento=fecha_base
    ))
    
    detalles_db[0].fecha_pago = fecha_base
    detalles_db[0].saldo_inicial = d2(monto_financiar)
    detalles_db[0].interes_capitalizado = d2(0)
    detalles_db[0].flujo_caja = d2(flujo_0)
    detalles_db[0].tea = None 
    detalles_db[0].tem = None 
    detalles_db[0].seguro_desgravamen = d2(0)
    detalles_db[0].cuota = d2(0)
    detalles_db[0].plazo_gracia = "-"

    for i in range(1, n_total + 1):
        saldo_anterior = saldo
        int_periodo = saldo_anterior * tem
        seguro_periodo = saldo_anterior * seguro_tasa
        gastos_periodicos = Decimal(str(payload.comision_periodica)) + Decimal(str(payload.portes)) + Decimal(str(payload.gastos_administracion))
        interes_cap = Decimal("0") 
        
        if i <= m_gracia:
            if payload.tipo_gracia == "Total":
                interes_cap = int_periodo
                amort_periodo = Decimal("0")
                seguro_pago = seguro_periodo
                cuota_base_pago = Decimal("0")
                cuota_t = cuota_base_pago + seguro_pago + gastos_periodicos
                saldo = saldo_anterior + interes_cap
            else: 
                amort_periodo = Decimal("0")
                seguro_pago = seguro_periodo
                cuota_base_pago = int_periodo
                cuota_t = cuota_base_pago + seguro_pago + gastos_periodicos
                saldo = saldo_anterior
                
            if i == m_gracia: 
                n_restantes = n_total - m_gracia
                factor_p = (tasa_para_factor * (1 + tasa_para_factor) ** n_restantes) / ((1 + tasa_para_factor) ** n_restantes - 1) if tasa_para_factor > 0 else (Decimal("1")/Decimal(str(n_restantes)))
                cuota_base = saldo * factor_p
        else:
            seguro_pago = seguro_periodo
            amort_periodo = cuota_base - int_periodo - seguro_pago
            cuota_base_pago = cuota_base
            cuota_t = cuota_base + gastos_periodicos
            saldo = saldo_anterior - amort_periodo
            
            if i == n_total and saldo != Decimal("0"):
                cuota_base_pago += saldo
                cuota_t += saldo
                amort_periodo += saldo
                saldo = Decimal("0")
        
        total_int += int_periodo
        total_seg += seguro_pago  
        flujos_caja.append(-float(cuota_t))

        detalle = SimulationDetail(
            numero_cuota=i,
            interes=d2(int_periodo),
            seguro=d2(seguro_pago),
            comision_periodica=d2(payload.comision_periodica),
            portes=d2(payload.portes),
            gastos_administracion=d2(payload.gastos_administracion),
            amortizacion=d2(amort_periodo), 
            cuota_total=d2(cuota_base_pago),
            saldo_final=d2(max(Decimal("0"), saldo)),
            fecha_vencimiento=date(fecha_base.year + (fecha_base.month + i - 1) // 12, (fecha_base.month + i - 1) % 12 + 1, min(fecha_base.day, 28))
        )
        
        detalle.saldo_inicio = d2(saldo_anterior)
        detalle.saldo_inicial = d2(saldo_anterior) 
        detalle.interes_capitalizado = d2(interes_cap)
        detalle.flujo_caja = d2(-cuota_t)
        detalle.fecha_pago = detalle.fecha_vencimiento 

        detalle.tea = tea * Decimal("100") 
        detalle.tem = tem * Decimal("100") 
        detalle.seguro_desgravamen = d2(seguro_periodo) 
        detalle.cuota = d2(cuota_t) 
        detalle.plazo_gracia = f"Gracia {payload.tipo_gracia}" if (payload.tipo_gracia != "Ninguno" and i <= m_gracia) else "Sin Gracia"

        detalles_db.append(detalle)

    max_cuota = max(d.cuota_total for d in detalles_db)
    
    ratio = (max_cuota / ifm * 100) if ifm > 0 else Decimal("0")
    
    if ifm > 0:
        limite_ratio = Decimal("40.00") if pv_pen <= Decimal("205000.00") else Decimal("50.00")
        if ratio > limite_ratio:
            raise HTTPException(
                status_code=400, 
                detail=(
                    f"Capacidad de pago excedida. El ratio cuota/ingreso es {float(ratio):.1f}%, "
                    f"superando el límite del {float(limite_ratio):.0f}% para este tipo de vivienda."
                )
            )

    TASA_DESCUENTO_ANUAL = Decimal("0.08")
    tasa_descuento_mensual = (1 + TASA_DESCUENTO_ANUAL) ** (Decimal("30") / Decimal("360")) - 1
    try:
        tir = npf.irr(flujos_caja)
        tcea = ((1 + tir) ** 12) - 1
        van = npf.npv(float(tasa_descuento_mensual), flujos_caja)
    except: tir, tcea, van = 0, 0, 0

    resumen_dict = {
        "rango_bbp": bono_info["rango"],
        "bono_bbp_base": float(bono_info["base"]),
        "bono_integrador_adicional": float(bono_info["integrador"]),
        "precio_neto": float(precio_neto),
        "monto_financiar": float(monto_financiar),
        "tasa_efectiva_anual": float(tea * 100),
        "tasa_efectiva_mensual": float(tem * 100),
        "tasa_descuento": float(TASA_DESCUENTO_ANUAL * 100),  
        "tasa_descuento_mensual": float(Decimal(str(tasa_descuento_mensual)) * 100),
        "factor_frances": float(factor),
        "cuota_base": float(cuota_base),
        "ratio_cuota_ingreso": float(ratio),
        "van": float(van),
        "tir": float(Decimal(str(tir)) * 100),
        "tcea": float(tcea * 100),
        "total_intereses": float(total_int),
        "total_pagado": float(sum(d.cuota_total for d in detalles_db if d.numero_cuota > 0)),
        "total_seguro": float(total_seg),
        "total_comisiones_periodicas": float(payload.comision_periodica * n_total),
        "total_portes_gastos_adm": float((payload.portes + payload.gastos_administracion) * n_total)
    }

    def _detalle_to_dict(d) -> dict:
        return {
            "numero_cuota":        d.numero_cuota,
            "fecha_vencimiento":   getattr(d, "fecha_vencimiento", None),
            "fecha_pago":          getattr(d, "fecha_pago", None),
            "tea":                 float(d.tea) if getattr(d, "tea", None) is not None else None,
            "tem":                 float(d.tem) if getattr(d, "tem", None) is not None else None,
            "plazo_gracia":        getattr(d, "plazo_gracia", "Sin Gracia"),
            "saldo_inicio":        float(d.saldo_inicio) if d.saldo_inicio is not None else None,
            "saldo_inicial":       float(getattr(d, "saldo_inicial", d.saldo_inicio)) if d.saldo_inicio is not None else None,
            "interes":             float(d.interes) if d.interes is not None else 0.0,
            "interes_capitalizado":float(getattr(d, "interes_capitalizado", 0) or 0),
            "amortizacion":        float(d.amortizacion) if d.amortizacion is not None else 0.0,
            "seguro":              float(d.seguro) if d.seguro is not None else 0.0,
            "seguro_desgravamen":  float(getattr(d, "seguro_desgravamen", d.seguro)) if d.seguro is not None else None,
            "comision_periodica":  float(d.comision_periodica) if getattr(d, "comision_periodica", None) is not None else 0.0,
            "portes":              float(d.portes) if getattr(d, "portes", None) is not None else 0.0,
            "gastos_administracion": float(d.gastos_administracion) if getattr(d, "gastos_administracion", None) is not None else 0.0,
            "cuota_total":         float(d.cuota_total) if d.cuota_total is not None else 0.0,
            "cuota":               float(d.cuota_total) if d.cuota_total is not None else 0.0,
            "saldo_final":         float(d.saldo_final) if d.saldo_final is not None else 0.0,
            "flujo_caja":          float(getattr(d, "flujo_caja", 0) or 0),
            "flujo":               float(getattr(d, "flujo_caja", 0) or 0),
        }

    try:
        if not save:
            return {
                "codigo_simulacion": None,
                "fecha_simulacion": date.today(),
                "fecha_inicio_prestamo": fecha_base,
                "cuota_inicial": float(payload.cuota_inicial),
                "gastos_iniciales": float(payload.gastos_iniciales),
                "coste_notarial": float(payload.coste_notarial),
                "coste_registral": float(payload.coste_registral),
                "tasacion": float(payload.tasacion),
                "comision_estudio": float(payload.comision_estudio),
                "comision_activacion": float(payload.comision_activacion),
                "tipo_bbp": payload.tipo_bbp,
                "bono_bbp": float(bono_total),
                "tipo_tasa": payload.tipo_tasa,
                "tasa_anual": float(payload.tasa_anual),
                "capitalizacion": payload.capitalizacion,
                "plazo_meses": payload.plazo_meses,
                "tipo_gracia": payload.tipo_gracia,
                "meses_gracia": payload.meses_gracia,
                "seguro_desgravamen": float(payload.seguro_desgravamen),
                "comision_periodica": float(payload.comision_periodica),
                "portes": float(payload.portes),
                "gastos_administracion": float(payload.gastos_administracion),
                "codigo_unidad": payload.codigo_unidad,
                "direccion_unidad": unit.direccion_unidad,
                "distrito_unidad": unit.distrito_unidad,
                "resumen": resumen_dict,
                "detalles": [_detalle_to_dict(d) for d in detalles_db]
            }

        new_sim = Simulation(
            fecha_inicio_prestamo=fecha_base,
            cuota_inicial=payload.cuota_inicial,
            coste_notarial=payload.coste_notarial,
            coste_registral=payload.coste_registral,
            tasacion=payload.tasacion,
            comision_estudio=payload.comision_estudio,
            comision_activacion=payload.comision_activacion,
            gastos_iniciales=payload.gastos_iniciales,
            tipo_bbp=payload.tipo_bbp, bono_bbp=bono_total,
            ifi_seleccionada=payload.ifi_seleccionada, tipo_tasa=payload.tipo_tasa,
            tasa_anual=payload.tasa_anual, plazo_meses=payload.plazo_meses,
            tipo_gracia=payload.tipo_gracia, meses_gracia=payload.meses_gracia,
            seguro_desgravamen=payload.seguro_desgravamen, 
            comision_periodica=payload.comision_periodica,
            portes=payload.portes,
            gastos_administracion=payload.gastos_administracion,
            codigo_unidad=payload.codigo_unidad, codigo_cliente=payload.codigo_cliente,
            codigo_prospecto=payload.codigo_prospecto, codigo_asesor=payload.codigo_asesor
        )
        db.add(new_sim); db.flush()

        resumen_obj = SimulationResult(
            codigo_simulacion=new_sim.codigo_simulacion,
            rango_bbp=bono_info["rango"], bono_bbp_base=bono_info["base"],
            bono_integrador_adicional=bono_info["integrador"],
            precio_neto=d2(precio_neto), monto_financiar=d2(monto_financiar),
            tasa_efectiva_anual=tea, tasa_efectiva_mensual=tem,
            factor_frances=factor, cuota_base=d2(cuota_base),
            ratio_cuota_ingreso=ratio, van=d2(van), tir=Decimal(str(tir)) * 100, tcea=Decimal(str(tcea * 100)),
            tasa_descuento=TASA_DESCUENTO_ANUAL * 100,
            tasa_descuento_mensual=Decimal(str(tasa_descuento_mensual)) * 100,
            total_intereses=d2(total_int),
            total_pagado=d2(sum(d.cuota_total for d in detalles_db if d.numero_cuota > 0)),
            total_seguro=d2(total_seg),
            total_comisiones_periodicas=d2(payload.comision_periodica * n_total),
            total_portes_gastos_adm=d2((payload.portes + payload.gastos_administracion) * n_total)
        )
        db.add(resumen_obj)
        for d in detalles_db: d.codigo_simulacion = new_sim.codigo_simulacion; db.add(d)
        
        db.commit(); db.refresh(new_sim)
        
        return {
            "codigo_simulacion": new_sim.codigo_simulacion,
            "fecha_simulacion": new_sim.fecha_simulacion,
            "fecha_inicio_prestamo": new_sim.fecha_inicio_prestamo,
            "cuota_inicial": float(new_sim.cuota_inicial),
            "gastos_iniciales": float(new_sim.gastos_iniciales),
            "coste_notarial": float(new_sim.coste_notarial),
            "coste_registral": float(new_sim.coste_registral),
            "tasacion": float(new_sim.tasacion),
            "comision_estudio": float(new_sim.comision_estudio),
            "comision_activacion": float(new_sim.comision_activacion),
            "tipo_bbp": new_sim.tipo_bbp,
            "bono_bbp": float(new_sim.bono_bbp),
            "tipo_tasa": new_sim.tipo_tasa,
            "tasa_anual": float(new_sim.tasa_anual),
            "comision_periodica": float(new_sim.comision_periodica),
            "portes": float(new_sim.portes),
            "gastos_administracion": float(new_sim.gastos_administracion),
            "capitalizacion": new_sim.capitalizacion,
            "plazo_meses": new_sim.plazo_meses,
            "tipo_gracia": new_sim.tipo_gracia,
            "meses_gracia": new_sim.meses_gracia,
            "seguro_desgravamen": float(new_sim.seguro_desgravamen),
            "codigo_unidad": new_sim.codigo_unidad,
            "direccion_unidad": unit.direccion_unidad,
            "distrito_unidad": unit.distrito_unidad,
            "resumen": resumen_dict,
            "detalles": detalles_db
        }

    except Exception as e:
        import traceback
        print(f"ERROR in run_simulation: {e}")
        print(traceback.format_exc())
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error al procesar simulación: {str(e)}")


@router.get("/", response_model=List[SimulationResponse])
def get_all_simulations(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    role = current_user.rol_rel.tipo_rol
    if role == "Administrador":
        return db.query(Simulation).all()
    elif role == "Asesor":
        return db.query(Simulation).filter(Simulation.codigo_asesor == current_user.codigo_usuario).all()
    else:
        raise HTTPException(status_code=403, detail="No tienes permiso.")


@router.get("/me", response_model=List[SimulationResponse])
def get_my_simulations(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    from sqlalchemy.orm import joinedload
    role = current_user.rol_rel.tipo_rol
    query = db.query(Simulation).options(
        joinedload(Simulation.unidad_rel),
        joinedload(Simulation.resumen)
    )
    
    if role == "Cliente":
        return query.filter(Simulation.codigo_cliente == current_user.codigo_usuario).all()
    elif role == "Asesor":
        return query.filter(Simulation.codigo_asesor == current_user.codigo_usuario).all()
    else:
        return query.all()


@router.get("/tem")
def convert_tea_to_tem(
    tea: float = Query(..., description="Tasa Efectiva Anual en porcentaje (ej: 10 para 10%)"),
    current_user: User = Depends(get_current_user)
):
    """
    Convierte una TEA (%) a TEM (%).
    Ejemplo: GET /simulator/tem?tea=10  →  { "tea": 10.0, "tem": 0.7974... }
    """
    tea_dec = Decimal(str(tea)) / Decimal("100")
    tem_dec = (1 + tea_dec) ** (Decimal("1") / Decimal("12")) - 1
    return {
        "tea": float(tea),
        "tem": float(tem_dec * 100)
    }


@router.get("/{codigo_simulacion}", response_model=SimulationResponse)
def get_simulation(
    codigo_simulacion: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    from sqlalchemy.orm import joinedload
    sim = db.query(Simulation)\
            .options(
                joinedload(Simulation.unidad_rel),
                joinedload(Simulation.resumen)
            )\
            .filter(Simulation.codigo_simulacion == codigo_simulacion)\
            .first()
    if not sim:
        raise HTTPException(status_code=404, detail="Simulación no encontrada")
    role = current_user.rol_rel.tipo_rol
    if role == "Cliente" and sim.codigo_cliente != current_user.codigo_usuario:
        raise HTTPException(status_code=403, detail="No tienes permiso para ver esta simulación.")
    return sim


@router.get("/client/{codigo_cliente}", response_model=List[SimulationResponse])
def get_simulations_by_client(
    codigo_cliente: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    role = current_user.rol_rel.tipo_rol
    if role == "Cliente" and current_user.codigo_usuario != codigo_cliente:
        raise HTTPException(status_code=403, detail="No puedes ver las simulaciones de otro cliente.")
    return db.query(Simulation).filter(Simulation.codigo_cliente == codigo_cliente).all()


@router.get("/advisor/{codigo_asesor}", response_model=List[SimulationResponse])
def get_simulations_by_advisor(
    codigo_asesor: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    role = current_user.rol_rel.tipo_rol
    if role == "Asesor" and current_user.codigo_usuario != codigo_asesor:
        raise HTTPException(status_code=403, detail="No puedes ver las simulaciones de otro asesor.")
    return db.query(Simulation).filter(Simulation.codigo_asesor == codigo_asesor).all()


@router.get("/{codigo_simulacion}/export/excel")
def export_simulation_excel(
    codigo_simulacion: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    sim = db.query(Simulation).filter(Simulation.codigo_simulacion == codigo_simulacion).first()
    if not (sim and sim.resumen):
        raise HTTPException(status_code=404, detail="Simulación o resultados no encontrados")
    role = current_user.rol_rel.tipo_rol
    if role == "Cliente" and sim.codigo_cliente != current_user.codigo_usuario:
        raise HTTPException(status_code=403, detail="No tienes permiso para exportar esta simulación.")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        res = sim.resumen
        detalles = sorted(sim.detalles, key=lambda x: x.numero_cuota)

        C_DARK, C_ORANGE, C_BLUE = "0F172A", "F97316", "3B82F6"
        C_LIGHT, C_ROW2, C_WHITE = "F8FAFC", "EFF6FF", "FFFFFF"
        C_GRAY, C_NAVY = "94A3B8", "1E293B"

        def fill(c): return PatternFill("solid", fgColor=c)
        def tb():
            s = Side(style="thin", color="E2E8F0")
            return Border(left=s, right=s, top=s, bottom=s)
        def hf(size=9, color=C_WHITE): return Font(name="Calibri", bold=True, color=color, size=size)
        def bf(bold=False, color=C_DARK, size=9): return Font(name="Calibri", bold=bold, color=color, size=size)
        def al(h="center", v="center"): return Alignment(horizontal=h, vertical=v)
        def mv(v): return float(v) if v is not None else 0.0

        wb = Workbook()
        # ─── Hoja 1: Resumen ────────────────────────────────────────────────────────
        ws = wb.active; ws.title = "Resumen"
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:D1")
        ws["A1"].value = "PROPEQUITY – Propuesta Financiera"
        ws["A1"].font = Font(name="Calibri", bold=True, color=C_WHITE, size=15)
        ws["A1"].fill, ws["A1"].alignment = fill(C_DARK), al()
        ws.row_dimensions[1].height = 34
        ws.merge_cells("A2:D2")
        ws["A2"].value = f"Simulación #{sim.codigo_simulacion}  ·  Unidad: {sim.codigo_unidad}  ·  {sim.fecha_simulacion.strftime('%d/%m/%Y') if sim.fecha_simulacion else ''}"
        ws["A2"].font = Font(name="Calibri", bold=True, color=C_ORANGE, size=9)
        ws["A2"].fill, ws["A2"].alignment = fill(C_NAVY), al()
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 6
        for col, (txt, bg) in enumerate([("Indicador", C_BLUE), ("Valor", C_BLUE), ("Parámetro", C_ORANGE), ("Detalle", C_ORANGE)], 1):
            c = ws.cell(row=4, column=col, value=txt)
            c.font, c.fill, c.alignment, c.border = hf(), fill(bg), al(), tb()
        ws.row_dimensions[4].height = 22
        indicadores = [
            ("Monto a Financiar",  f"S/ {mv(res.monto_financiar):,.2f}"),
            ("Cuota Base",         f"S/ {mv(res.cuota_base):,.2f}"),
            ("TEA",                f"{mv(res.tasa_efectiva_anual)*100:.4f}%"),
            ("TEM",                f"{mv(res.tasa_efectiva_mensual)*100:.6f}%"),
            ("TCEA",               f"{mv(res.tcea):.2f}%"),
            ("VAN",                f"S/ {mv(res.van):,.2f}"),
            ("TIR (mensual)",      f"{mv(res.tir):.5f}%"),
            ("Intereses",          f"S/ {mv(res.total_intereses):,.2f}"),
            ("Comisiones",         f"S/ {mv(res.total_comisiones_periodicas):,.2f}"),
            ("Portes / Otros",     f"S/ {mv(res.total_portes_gastos_adm):,.2f}"),
            ("Seguros",            f"S/ {mv(res.total_seguro):,.2f}"),
            ("Tasa Descuento (M)", f"{mv(res.tasa_descuento_mensual):.5f}%"),
            ("Total Pagado",       f"S/ {mv(res.total_pagado):,.2f}"),
        ]
        parametros = [
            ("IFI",           sim.ifi_seleccionada or "Genérico"),
            ("Tipo de Tasa",  sim.tipo_tasa),
            ("Tasa Anual",    f"{float(sim.tasa_anual):.2f}%"),
            ("Plazo",         f"{sim.plazo_meses} meses"),
            ("Tipo BBP",      sim.tipo_bbp),
            ("Bono BBP",      f"S/ {mv(sim.bono_bbp):,.2f}"),
            ("Tipo Gracia",   sim.tipo_gracia),
            ("Meses Gracia",  str(sim.meses_gracia)),
            ("Seg. Desgrav.", f"{float(sim.seguro_desgravamen):.4f}%"),
            ("Rango BBP",     res.rango_bbp or "—"),
        ]
        for i, ((ind, val), (par, det)) in enumerate(zip(indicadores, parametros)):
            row = 5 + i; bg = C_LIGHT if i % 2 == 0 else C_ROW2
            ws.row_dimensions[row].height = 18
            for col, (txt, bold, h) in enumerate([(ind, False, "left"), (val, True, "right"), (par, False, "left"), (det, True, "right")], 1):
                c = ws.cell(row=row, column=col, value=txt)
                c.font, c.fill, c.border, c.alignment = bf(bold=bold, color=C_DARK if bold else C_GRAY), fill(bg), tb(), al(h=h)
        for col, w in zip("ABCD", [24, 22, 22, 20]): ws.column_dimensions[col].width = w

        # ─── Hoja 2: Cronograma ───────────────────────────────────────────────────────
        ws2 = wb.create_sheet("Cronograma de Pagos")
        ws2.sheet_view.showGridLines = False
        ws2.merge_cells("A1:H1")
        ws2["A1"].value = "CRONOGRAMA DE PAGOS – PropEquity"
        ws2["A1"].font = Font(name="Calibri", bold=True, color=C_WHITE, size=13)
        ws2["A1"].fill, ws2["A1"].alignment = fill(C_DARK), al()
        ws2.row_dimensions[1].height = 30
        ws2.merge_cells("A2:K2")
        ws2["A2"].value = f"Sim #{sim.codigo_simulacion}  |  S/ {mv(res.monto_financiar):,.2f}  |  {sim.plazo_meses} m  |  TEA {mv(res.tasa_efectiva_anual)*100:.2f}%  |  TCEA {mv(res.tcea):.2f}%"
        ws2["A2"].font = Font(name="Calibri", bold=True, color=C_ORANGE, size=8)
        ws2["A2"].fill, ws2["A2"].alignment = fill(C_NAVY), al()
        ws2.row_dimensions[2].height = 16; ws2.row_dimensions[3].height = 6
        hdrs = [("N°", 7), ("Fecha Pago", 13), ("Saldo Inicial", 16), ("Interés", 14),
                ("Amortización", 14), ("Seg. Desgrav.", 14), ("Comisión", 12), ("Portes", 12), ("Gastos Adm", 12), ("Flujo", 14), ("Saldo Final", 16)]
        ws2.row_dimensions[4].height = 22
        for col, (hdr, w) in enumerate(hdrs, 1):
            c = ws2.cell(row=4, column=col, value=hdr)
            c.font, c.fill, c.alignment, c.border = hf(), fill(C_ORANGE), al(), tb()
            ws2.column_dimensions[get_column_letter(col)].width = w
        for i, d in enumerate(detalles):
            row = 5 + i; bg = "DBEAFE" if d.numero_cuota == 0 else (C_LIGHT if i % 2 == 0 else C_ROW2)
            ws2.row_dimensions[row].height = 15
            row_vals = [d.numero_cuota,
                        d.fecha_vencimiento.strftime("%d/%m/%Y") if d.fecha_vencimiento else "—",
                        mv(d.saldo_inicio), mv(d.interes), mv(d.amortizacion),
                        mv(d.seguro), mv(d.comision_periodica), mv(d.portes), mv(d.gastos_administracion),
                        mv(d.cuota_total), mv(d.saldo_final)]
            for col, val in enumerate(row_vals, 1):
                c = ws2.cell(row=row, column=col, value=val)
                is_cuota = col == 10; is_num = col >= 3
                c.font = bf(bold=is_cuota, color=C_ORANGE if is_cuota else C_DARK)
                c.fill, c.border = fill(bg), tb()
                c.alignment = al(h="right") if col >= 3 else al()
                if is_num: c.number_format = "#,##0.00"
        tr = 5 + len(detalles); ws2.row_dimensions[tr].height = 20
        ws2.merge_cells(f"A{tr}:B{tr}")
        c = ws2[f"A{tr}"]; c.value = "TOTALES"; c.font = hf(); c.fill = fill(C_DARK); c.alignment = al(); c.border = tb()
        total_map = {4: mv(res.total_intereses), 6: mv(res.total_seguro), 7: mv(res.total_comisiones_periodicas), 8: mv(res.total_portes_gastos_adm), 10: mv(res.total_pagado)}
        for col in range(3, 12):
            c = ws2.cell(row=tr, column=col, value=total_map.get(col))
            c.font, c.fill, c.border, c.alignment, c.number_format = hf(color=C_ORANGE), fill(C_DARK), tb(), al(h="right"), "#,##0.00"

        output = io.BytesIO()
        wb.save(output); output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=PropEquity_Sim_{codigo_simulacion}.xlsx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando Excel: {str(e)}")


@router.get("/{codigo_simulacion}/export/pdf")
def export_simulation_pdf(
    codigo_simulacion: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    sim = db.query(Simulation).filter(Simulation.codigo_simulacion == codigo_simulacion).first()
    if not (sim and sim.resumen):
        raise HTTPException(status_code=404, detail="Simulación no encontrada")
    role = current_user.rol_rel.tipo_rol
    if role == "Cliente" and sim.codigo_cliente != current_user.codigo_usuario:
        raise HTTPException(status_code=403, detail="No tienes permiso para exportar esta simulación.")

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()
    res = sim.resumen

    def mv(v): return float(v) if v is not None else 0.0

    elements.append(Paragraph(f"PropEquity - Propuesta Financiera #{sim.codigo_simulacion}", styles['Title']))
    elements.append(Spacer(1, 12))
    summary = f"""
    <b>Unidad:</b> {sim.unidad_rel.direccion_unidad}<br/>
    <b>TCEA:</b> {float(res.tcea):.5f}% | <b>TIR Mes:</b> {float(res.tir):.5f}% | <b>VAN:</b> S/ {res.van:,.2f}<br/>
    <b>Monto Financiar:</b> S/ {res.monto_financiar:,.2f} | <b>Total Pagado:</b> S/ {res.total_pagado:,.2f}<br/>
    <b>Intereses Totales:</b> S/ {res.total_intereses:,.2f} | <b>Seguros Totales:</b> S/ {res.total_seguro:,.2f}<br/>
    <b>Comisiones Totales:</b> S/ {res.total_comisiones_periodicas:,.2f} | <b>Portes/Gastos:</b> S/ {res.total_portes_gastos_adm:,.2f}<br/>
    <b>Tasa Descuento (Mensual):</b> {float(res.tasa_descuento_mensual):.5f}%
    """
    elements.append(Paragraph(summary, styles['Normal']))
    elements.append(Spacer(1, 20))

    data = [["Cuota", "Fecha", "Interés", "Amortiz.", "Seguro", "Comis.", "Porte", "Gasto", "Flujo", "Saldo"]]
    for d in sim.detalles:
        data.append([d.numero_cuota, d.fecha_vencimiento.strftime("%d/%m/%Y"),
                     f"{d.interes:,.2f}", f"{d.amortizacion:,.2f}", f"{d.seguro:,.2f}", 
                     f"{d.comision_periodica:,.2f}", f"{d.portes:,.2f}", f"{d.gastos_administracion:,.2f}",
                     f"{d.cuota_total:,.2f}", f"{d.saldo_final:,.2f}"])

    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.grey), ('GRID', (0, 0), (-1, -1), 1, colors.black)]))
    elements.append(t)
    doc.build(elements)
    output.seek(0)
    return StreamingResponse(output, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=Propuesta_{codigo_simulacion}.pdf"})